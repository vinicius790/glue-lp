#!/usr/bin/env python3
"""Build GLUE-LP scientific DOCX/PDF from the bulk of docs/ARTIGO.md.

Incorporates headings, paragraphs, tables and fenced code from ARTIGO.md
(not a thin summary). Figures fig1–fig9 inserted at first Results mention.
Numbers are not invented here — they live in ARTIGO.md / experiments/*.json.
"""
from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ARTIGO = ROOT / "docs" / "ARTIGO.md"
FIG = ROOT / "figures"
OUT_ROOT = ROOT / "GLUE-LP-artigo-cientifico.docx"
OUT_DOCS = ROOT / "docs" / "GLUE-LP-artigo-cientifico.docx"
PDF_ROOT = ROOT / "GLUE-LP-artigo-cientifico.pdf"
PDF_DOCS = ROOT / "docs" / "GLUE-LP-artigo-cientifico.pdf"

FIGURES = [
    ("fig1_pipeline.png", "Figura 1. Pipeline GLUE-LP: dados → split → assert → encoder → mesmo Q."),
    ("fig2_protocolo.png", "Figura 2. Protocolo valid vs leaky (E_mp)."),
    ("fig3_techtree.png", "Figura 3. Tech tree sintético (domínio ilustrativo)."),
    ("fig4_cora_auc.png", "Figura 4. AUC Cora — série 1 (pré-correção ES)."),
    ("fig5_cora_citeseer.png", "Figura 5. Cora × Citeseer — série 2 (pré-correção ES)."),
    ("fig6_seeds.png", "Figura 6. Dispersão por seed (valid vs leaky)."),
    ("fig7_tres_corpora.png", "Figura 7. Três corpora LINQS (pré-correção ES)."),
    ("fig8_arquitetura.png", "Figura 8. Arquitetura de módulos do artefato glue_lp."),
    ("fig9_escada_vazamento.png", "Figura 9. Escada de vazamento L1 / L3 / L4 (sintético)."),
]


def set_run_font(run, *, size=11, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph):
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r = paragraph.add_run()
    r._r.append(fld1)
    r._r.append(instr)
    r._r.append(fld2)
    set_run_font(r, size=8)


def P(doc, text, *, italic=False, bold=False, size=11, center=False, space_after=6, first_indent=True):
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.15
    if first_indent and not center:
        pf.first_line_indent = Cm(0.75)
    # light markdown: **bold** and *italic*
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, italic=italic)
        elif part.startswith("*") and part.endswith("*"):
            run = para.add_run(part[1:-1])
            set_run_font(run, size=size, bold=bold, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = para.add_run(part[1:-1])
            set_run_font(run, size=size - 1, bold=bold, italic=italic, name="Courier New")
        else:
            run = para.add_run(part)
            set_run_font(run, size=size, bold=bold, italic=italic)
    return para


def heading(doc, text, level=1):
    # strip markdown # prefixes
    text = re.sub(r"^#+\s*", "", text).strip()
    h = doc.add_heading(text, level=min(level, 3))
    for run in h.runs:
        set_run_font(run, size={1: 14, 2: 12, 3: 11}.get(level, 11), bold=True)
    h.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.line_spacing = 1.15
    return h


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            val = row[j] if j < len(row) else ""
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=8, bold=(i == 0))
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_figure(doc, name, caption):
    path = FIG / name
    if not path.exists():
        P(doc, f"[figura ausente: {name}]", italic=True, first_indent=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(5.8))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, size=9, italic=True)
    cap.paragraph_format.space_after = Pt(8)


def parse_md_table(lines):
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def strip_tex(s: str) -> str:
    s = re.sub(r"\$\$([^$]+)\$\$", r"\1", s)
    s = re.sub(r"\$([^$]+)\$", r"\1", s)
    s = s.replace("\\(", "").replace("\\)", "")
    s = s.replace("\\[", "").replace("\\]", "")
    s = re.sub(r"\\(?:text|mathrm|operatorname)\{([^}]+)\}", r"\1", s)
    s = re.sub(r"\\(?:frac)\{([^}]+)\}\{([^}]+)\}", r"(\1)/(\2)", s)
    s = s.replace("\\emptyset", "∅").replace("\\cap", "∩").replace("\\cup", "∪")
    s = s.replace("\\subseteq", "⊆").replace("\\in", "∈").replace("\\notin", "∉")
    s = s.replace("\\langle", "⟨").replace("\\rangle", "⟩")
    s = s.replace("\\cdot", "·").replace("\\times", "×")
    s = s.replace("\\pm", "±").replace("\\approx", "≈")
    s = s.replace("\\theta", "θ").replace("\\hat", "")
    s = re.sub(r"\{([^{}]+)\}", r"\1", s)
    s = s.replace("\\", "")
    return s


def build():
    raw = ARTIGO.read_text(encoding="utf-8")
    lines = raw.splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run("GLUE-LP · Graph Link-evaluation Under Exclusion · manuscrito Design Science")
    set_run_font(r, size=8, italic=True)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = fp.add_run("glue_lp v0.5.0 · apenas experiments/*.json · p. ")
    set_run_font(rr, size=8)
    add_page_number(fp)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.15

    figures_done = False
    i = 0
    n = len(lines)
    para_buf: list[str] = []

    def flush_para():
        nonlocal para_buf
        if not para_buf:
            return
        text = " ".join(para_buf).strip()
        para_buf = []
        if not text:
            return
        text = strip_tex(text)
        # blockquotes
        if text.startswith(">"):
            text = text.lstrip("> ").strip()
            P(doc, text, italic=True, size=10, first_indent=False, space_after=6)
        else:
            P(doc, text, size=11, space_after=6)

    while i < n:
        line = lines[i]
        # code fence
        if line.strip().startswith("```"):
            flush_para()
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # closing fence
            block = "\n".join(code)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(block)
            set_run_font(run, size=8, name="Courier New")
            continue

        # table
        if line.strip().startswith("|") and i + 1 < n and re.search(r"\|?\s*:?--+:?", lines[i + 1]):
            flush_para()
            tbl_lines = []
            while i < n and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            add_table(doc, parse_md_table(tbl_lines))
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush_para()
            level = len(m.group(1))
            title = strip_tex(m.group(2))
            # Skip duplicating H1 title handling — render all
            if level == 1 and title.startswith("GLUE-LP:"):
                P(doc, title, bold=True, size=16, center=True, space_after=8, first_indent=False)
            else:
                heading(doc, title, level=min(level, 3))
            # Insert figures once we hit Resultados
            if (not figures_done) and ("Resultados" in title or title.startswith("7.")):
                for fname, cap in FIGURES:
                    add_figure(doc, fname, cap)
                figures_done = True
            i += 1
            continue

        # horizontal rule
        if re.match(r"^---+$", line.strip()):
            flush_para()
            i += 1
            continue

        # blank
        if not line.strip():
            flush_para()
            i += 1
            continue

        # list items as paragraphs
        if re.match(r"^(\s*[-*]|\s*\d+\.)\s+", line):
            flush_para()
            text = strip_tex(re.sub(r"^(\s*[-*]|\s*\d+\.)\s+", "• ", line).strip())
            P(doc, text, size=11, first_indent=False, space_after=3)
            i += 1
            continue

        para_buf.append(line.strip())
        i += 1

    flush_para()

    if not figures_done:
        heading(doc, "Figuras", level=1)
        for fname, cap in FIGURES:
            add_figure(doc, fname, cap)

    OUT_ROOT.parent.mkdir(parents=True, exist_ok=True)
    OUT_DOCS.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_ROOT))
    shutil.copy2(OUT_ROOT, OUT_DOCS)
    print(f"DOCX written: {OUT_ROOT}")
    print(f"DOCX copy:    {OUT_DOCS}")

    # PDF via LibreOffice
    for out_pdf_dir, label in [(ROOT, "root"), (ROOT / "docs", "docs")]:
        cmd = [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_pdf_dir),
            str(OUT_ROOT if label == "root" else OUT_DOCS),
        ]
        print("Running:", " ".join(cmd))
        subprocess.run(cmd, check=True, capture_output=True, text=True)

    for pdf in (PDF_ROOT, PDF_DOCS):
        print(f"PDF: {pdf} exists={pdf.exists()} size={pdf.stat().st_size if pdf.exists() else 0}")


if __name__ == "__main__":
    build()
